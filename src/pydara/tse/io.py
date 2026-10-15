"""Explicit, synchronous, non-executable table archives.

CSV uses a .pydara.json sidecar. Excel and Parquet embed structural metadata.
SQLite holds named snapshots in a dedicated database, not the live cache.
Explicit failures raise: only automatic persistence is best-effort.
"""
from __future__ import annotations
from contextlib import contextmanager
from datetime import UTC, datetime
from hashlib import sha256
import csv
import io
import os
from pathlib import Path
import sqlite3
import tempfile
import uuid
import zipfile
import zlib

import pandas as pd
from pydara.tse._core.models import WriteReceipt
from pydara.tse._core.exceptions import (PyDaraError, FileExistsDataError, InvalidArgumentError,
    OptionalDependencyError, StorageBusyError, StorageError)
from .storage.codec import (NULL, dumps, loads, freeze, thaw, FrozenFrame,
                            encode_cell, string_wire, string_unwire)

MAX_BYTES=512*1024*1024
FORMATS={'.csv':'csv','.xlsx':'excel','.parquet':'parquet','.sqlite':'sqlite','.sqlite3':'sqlite','.db':'sqlite'}


def dependency(module, extra):
    """Perform the dependency operation.
    
    Args:
        module (Any): Input for the dependency operation.
        extra (Any): Input for the dependency operation.
    
    Returns:
        Implementation-defined result or generator output; see the calling API.
    
    Raises:
        OptionalDependencyError: When the corresponding validation or operation fails."""
    import importlib
    try: return importlib.import_module(module)
    except ImportError as exc:
        raise OptionalDependencyError(
            f'{module} is required for this format. Install pydara[{extra}] '
            f'or install {module} into the active environment.') from exc


def file_format(path, format):
    """Perform the file format operation.
    
    Args:
        path (Any): Local file or database path.
        format (Any): Requested output format.
    
    Returns:
        Implementation-defined result or generator output; see the calling API.
    
    Raises:
        InvalidArgumentError: When the corresponding validation or operation fails."""
    if not isinstance(path, (str, Path)) or isinstance(path, str) and not path.strip():
        raise InvalidArgumentError('path must be a nonempty string or pathlib.Path.')
    if format is not None and not isinstance(format, str):
        raise InvalidArgumentError('format must be csv, excel, parquet, sqlite, or None.')
    fmt=FORMATS.get(Path(path).suffix.lower()) if format is None else format.lower()
    if fmt=='xlsx': fmt='excel'
    if fmt not in set(FORMATS.values()):
        raise InvalidArgumentError('Supported formats: csv, excel (.xlsx), parquet, sqlite.')
    return fmt


@contextmanager
def path_lock(path):
    """Perform the path lock operation.
    
    Args:
        path (Any): Local file or database path.
    
    Returns:
        Implementation-defined result or generator output; see the calling API.
    
    Raises:
        StorageBusyError: When the corresponding validation or operation fails."""
    lock=path.with_name(path.name+'.pydara.lock')
    try: fd=os.open(lock,os.O_CREAT|os.O_EXCL|os.O_WRONLY,0o600)
    except FileExistsError as exc:
        raise StorageBusyError(f'Another export owns {lock.name}; stale locks require manual review.') from exc
    try:
        with os.fdopen(fd, 'wb') as handle:
            handle.write(str(os.getpid()).encode())
        yield
    finally:
        lock.unlink(missing_ok=True)


def temp_path(path):
    """Perform the temp path operation.
    
    Args:
        path (Any): Local file or database path.
    
    Returns:
        Implementation-defined result or generator output; see the calling API."""
    fd,name=tempfile.mkstemp(prefix='.'+path.name+'.',suffix='.tmp',dir=path.parent)
    os.close(fd)
    return Path(name)


def sync_file(path):
    # Windows' os.fsync requires a writable file descriptor.
    # The temporary export is already fully written and closed.
    """Synchronize file.
    
    Args:
        path (Any): Local file or database path."""
    with open(path,'r+b') as f: os.fsync(f.fileno())


def _native(text,spec):
    """Process internal native data.
    
    Args:
        text (Any): Input for the _native operation.
        spec (Any): Request or schema specification.
    
    Returns:
        Implementation-defined result or generator output; see the calling API."""
    if text==NULL: return None
    kind=spec['kind']
    if kind=='int': return int(text)
    if kind=='float': return float(text)
    if kind=='bool': return text=='1'
    return text


def _wire(value,spec):
    """Process internal wire data.
    
    Args:
        value (Any): Input for the _wire operation.
        spec (Any): Request or schema specification.
    
    Returns:
        Implementation-defined result or generator output; see the calling API."""
    if value is None: return NULL
    kind=spec['kind']
    if kind=='float': return repr(float(value))
    if kind=='int': return str(int(value))
    if kind=='bool': return '1' if bool(value) else '0'
    return str(value)


def _write_xlsx(frozen, path):
    """Write xlsx.
    
    Args:
        frozen (Any): Input for the _write_xlsx operation.
        path (Any): Local file or database path.
    
    Raises:
        StorageError: When the corresponding validation or operation fails."""
    ox=dependency('openpyxl','excel')
    from openpyxl.styles import Font, PatternFill, Alignment
    if len(frozen.rows)>1048575 or len(frozen.manifest['columns'])>16384:
        raise StorageError('Excel sheet limits exceeded; use Parquet or SQLite.')
    wb=ox.Workbook(); ws=wb.active;ws.title='data';ws.freeze_panes='A2';ws.sheet_view.showGridLines=False
    specs=frozen.manifest['specs']
    def cell_value(c,value):
        """Perform the cell value operation.
        
        Args:
            c (Any): Input for the cell_value operation.
            value (Any): Input for the cell_value operation.
        
        Raises:
            StorageError: When the corresponding validation or operation fails."""
        if isinstance(value,str):
            if len(value)>32767: raise StorageError('Excel cell text exceeds 32767 characters.')
            c.value=value;c.data_type='s'
        else: c.value=value
    try:
        for j,label in enumerate(frozen.manifest['columns'],1):
            c=ws.cell(1,j);cell_value(c,label)
            c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='163A4A')
            ws.column_dimensions[c.column_letter].width=min(40,max(16,len(label)+2))
        for i,row in enumerate(frozen.rows,2):
            for j,(text,spec) in enumerate(zip(row,specs),1):
                value=text;kind=spec['kind']
                if text!=NULL and kind in ('int','float','bool'):
                    native=_native(text,spec)
                    # Excel only reliably carries 15 significant decimal digits.
                    if kind=='bool' or (kind=='int' and len(text.lstrip('-'))<=15): value=native
                    elif kind=='float' and text not in ('inf','-inf','nan','-0.0'):
                        if float(format(native,'.15g'))==native: value=native
                c=ws.cell(i,j);cell_value(c,value)
                c.font=Font(color='008000')
                if isinstance(value,(int,float)) and not isinstance(value,bool): c.number_format='#,##0.########'
        ws.auto_filter.ref=ws.dimensions
        meta=wb.create_sheet('_pydara_meta')
        meta.sheet_state='hidden'
        text=dumps(frozen.manifest)
        for i,start in enumerate(range(0,len(text),30000),1):
            cell_value(meta.cell(i,1),text[start:start+30000])
        wb.save(path)
    finally: wb.close()


def _read_xlsx(path):
    """Read xlsx.
    
    Args:
        path (Any): Local file or database path.
    
    Returns:
        Implementation-defined result or generator output; see the calling API.
    
    Raises:
        StorageError: When the corresponding validation or operation fails."""
    ox=dependency('openpyxl','excel')
    with zipfile.ZipFile(path) as z:
        if sum(i.file_size for i in z.infolist())>MAX_BYTES:
            raise StorageError('Uncompressed workbook exceeds safety limit.')
    wb=ox.load_workbook(path,read_only=True,data_only=False,keep_links=False)
    try:
        metadata_sheet = '_pydara_meta' if '_pydara_meta' in wb else '_dara_meta'
        if metadata_sheet not in wb or 'data' not in wb:
            raise StorageError('Not a PyDara Excel archive; use pandas.read_excel for arbitrary workbooks.')
        manifest=loads(''.join(str(r[0]) for r in wb[metadata_sheet].iter_rows(values_only=True)))
        iterator=wb['data'].iter_rows()
        header=next(iterator)
        if [c.value for c in header]!=manifest['columns']:
            raise StorageError('Excel column headers do not match the manifest.')
        rows=[]
        for row in iterator:
            if any(c.data_type=='f' for c in row):
                raise StorageError('Formulas are not permitted in a PyDara data archive.')
            # Numeric native cells are converted back through their declared logical type.
            rows.append([_wire(c.value,s) if not isinstance(c.value,str) else c.value
                         for c,s in zip(row,manifest['specs'])])
        return manifest,rows
    finally: wb.close()


def _write_parquet(frozen,path):
    """Write parquet.
    
    Args:
        frozen (Any): Input for the _write_parquet operation.
        path (Any): Local file or database path."""
    pa=dependency('pyarrow','parquet');pq=dependency('pyarrow.parquet','parquet')
    arrays={}
    for i,(col,spec) in enumerate(zip(frozen.manifest['columns'],frozen.manifest['specs'])):
        kind=spec['kind']
        typ=pa.float64() if kind=='float' else pa.bool_() if kind=='bool' else (
            pa.uint64() if spec['dtype'].lower().startswith('uint') else pa.int64()) if kind=='int' else pa.string()
        arrays[col]=pa.array([_native(r[i],spec) for r in frozen.rows],type=typ)
    table=pa.table(arrays).replace_schema_metadata({b'pydara':dumps(frozen.manifest).encode()})
    pq.write_table(table,path,compression='zstd')


def _read_parquet(path):
    """Read parquet.
    
    Args:
        path (Any): Local file or database path.
    
    Returns:
        Implementation-defined result or generator output; see the calling API.
    
    Raises:
        StorageError: When the corresponding validation or operation fails."""
    pq=dependency('pyarrow.parquet','parquet')
    table=pq.read_table(path)
    metadata=table.schema.metadata or {}
    metadata_key = next((key for key in (b'pydara', b'tse_dara', b'dara') if key in metadata), b'pydara')
    if metadata_key not in metadata:
        raise StorageError('No PyDara manifest. Use pandas.read_parquet for arbitrary files.')
    manifest=loads(metadata[metadata_key].decode())
    if table.column_names!=manifest['columns']: raise StorageError('Parquet schema mismatch.')
    cols=[table[c].to_pylist() for c in table.column_names]
    rows=[[_wire(v,s) for v,s in zip(row,manifest['specs'])] for row in zip(*cols)]
    return manifest,rows


def _sqlite_frame(path, dataset, frozen=None, overwrite=False):
    """Process internal sqlite frame data.
    
    Args:
        path (Any): Local file or database path.
        dataset (Any): Input for the _sqlite_frame operation.
        frozen (Any): Input for the _sqlite_frame operation.
        overwrite (Any): Input for the _sqlite_frame operation.
    
    Returns:
        Implementation-defined result or generator output; see the calling API.
    
    Raises:
        FileExistsDataError: When the corresponding validation or operation fails.
        InvalidArgumentError: When the corresponding validation or operation fails.
        StorageError: When the corresponding validation or operation fails."""
    if not isinstance(dataset,str) or not dataset or len(dataset)>200:
        raise InvalidArgumentError('dataset must be a nonempty name of at most 200 characters.')
    if frozen is None:
        con=sqlite3.connect(path.absolute().as_uri()+'?mode=ro',uri=True,timeout=1)
    else: con=sqlite3.connect(str(path),timeout=1)
    try:
        tables={r[0] for r in con.execute("SELECT name FROM sqlite_master WHERE type='table'")}
        if tables and 'dara_frame_archives' not in tables:
            raise StorageError('Not a PyDara snapshot database. Do not export over the live response cache.')
        if frozen is None:
            if 'dara_frame_archives' not in tables: raise StorageError('No snapshot tables.')
            row=con.execute('SELECT manifest,content FROM dara_frame_archives WHERE dataset=?',(dataset,)).fetchone()
            if row is None: raise StorageError(f'Unknown stored dataset: {dataset!r}.')
            dec=zlib.decompressobj(); raw=dec.decompress(row[1],MAX_BYTES+1)
            if len(raw)>MAX_BYTES or dec.unconsumed_tail or not dec.eof:
                raise StorageError('Snapshot payload exceeds limits or is incomplete.')
            return loads(row[0]), loads(raw.decode())
        con.execute('PRAGMA synchronous=FULL')
        con.execute('BEGIN IMMEDIATE')
        con.execute('CREATE TABLE IF NOT EXISTS dara_frame_archives ('
                    'dataset TEXT PRIMARY KEY,manifest TEXT NOT NULL,content BLOB NOT NULL)')
        if not overwrite and con.execute('SELECT 1 FROM dara_frame_archives WHERE dataset=?',(dataset,)).fetchone():
            raise FileExistsDataError(f'Dataset {dataset!r} already exists.')
        blob=zlib.compress(dumps(frozen.rows).encode())
        con.execute('INSERT INTO dara_frame_archives VALUES(?,?,?) ON CONFLICT(dataset) DO UPDATE SET '
                    'manifest=excluded.manifest,content=excluded.content',
                    (dataset,dumps(frozen.manifest),blob))
        con.commit()
    except BaseException:
        con.rollback();raise
    finally: con.close()


def save(df: pd.DataFrame,path: str | Path,*,format: str | None=None,dataset: str='user_data',
         metadata: bool=True,overwrite: bool=False,safe_text: bool=True) -> WriteReceipt:
    """Save a self-describing snapshot; return only after successful commit.

    ``metadata=False`` omits df.attrs, not structural type/checksum information.
    MultiIndex/unsupported object values are rejected. No network requests.
    CSV + sidecar replace is not one filesystem transaction: a crash between
    replacements is detected by load's SHA256 check, never silently accepted.
    """
    fmt=file_format(path,format);path=Path(path).absolute()
    if not isinstance(overwrite,bool):raise InvalidArgumentError('overwrite must be bool.')
    frozen=freeze(df,safe_text=safe_text,metadata=metadata)
    accepted=datetime.now(UTC);temporary=[]
    try:
        path.parent.mkdir(parents=True,exist_ok=True)
        with path_lock(path):
            if fmt=='sqlite': _sqlite_frame(path,dataset,frozen,overwrite)
            else:
                side=path.with_name(path.name+'.pydara.json')
                if not overwrite and (path.exists() or (fmt=='csv' and side.exists())):
                    raise FileExistsDataError(f'File already exists: {path.name}.')
                tmp=temp_path(path);temporary.append(tmp)
                if fmt=='csv':
                    with open(tmp,'w',encoding='utf-8-sig',newline='') as f:
                        writer=csv.writer(f);writer.writerow([string_wire(c,safe_text) for c in frozen.manifest['columns']]);writer.writerows(frozen.rows)
                    frozen.manifest['file_sha256']=sha256(tmp.read_bytes()).hexdigest()
                    tm=temp_path(side);temporary.append(tm);tm.write_text(dumps(frozen.manifest),encoding='utf-8')
                    sync_file(tm)
                elif fmt=='excel': _write_xlsx(frozen,tmp)
                else: _write_parquet(frozen,tmp)
                sync_file(tmp);os.replace(tmp,path)
                if fmt=='csv': os.replace(tm,side)
        return WriteReceipt(receipt_id=uuid.uuid4().hex,request_id=None,state='persisted',
            accepted_at=accepted,committed_at=datetime.now(UTC),rows=len(df),payload_bytes=path.stat().st_size,
            reason=None,destination=str(path),content_hash=frozen.manifest['matrix_sha256'])
    except PyDaraError: raise
    except Exception as exc:
        raise StorageError(f'Export failed ({type(exc).__name__}): {exc}') from exc
    finally:
        for p in temporary: p.unlink(missing_ok=True)


def load(path: str | Path,*,format: str | None=None,dataset: str='user_data',
         schema=None,verify: bool=True) -> pd.DataFrame:
    """Read a PyDara archive and verify its structure/content. This is not a signature.

    For arbitrary externally-created files use pandas readers, then validate.
    ``verify=False`` allows deliberate value edits, not code execution or a
    different schema. CSV still requires its sidecar for dtype restoration.
    """
    if schema is not None: raise InvalidArgumentError('schema override is reserved; archives carry their own schema.')
    if not isinstance(verify,bool):raise InvalidArgumentError('verify must be bool.')
    fmt=file_format(path,format);path=Path(path).absolute()
    try:
        if not path.is_file():raise StorageError('Archive file not found.')
        if path.stat().st_size>MAX_BYTES and fmt!='sqlite':raise StorageError('Archive exceeds 512 MiB safety limit.')
        if fmt=='csv':
            side=path.with_name(path.name+'.pydara.json')
            if not side.is_file():
                side=path.with_name(path.name+'.dara.json')
            if not side.is_file():raise StorageError('CSV sidecar is missing; keep .csv and .csv.pydara.json together.')
            if side.stat().st_size>32*1024*1024:raise StorageError('Manifest too large.')
            manifest=loads(side.read_text(encoding='utf-8'))
            if verify and sha256(path.read_bytes()).hexdigest()!=manifest.get('file_sha256'):
                raise StorageError('CSV checksum differs from its sidecar; edited or incomplete export.')
            with open(path,encoding='utf-8-sig',newline='') as f:
                reader=csv.reader(f);header=[string_unwire(c,manifest['safe_text']) for c in next(reader)]
                if header!=manifest['columns']:raise StorageError('CSV header mismatch.')
                rows=list(reader)
        elif fmt=='excel':manifest,rows=_read_xlsx(path)
        elif fmt=='parquet':manifest,rows=_read_parquet(path)
        else:manifest,rows=_sqlite_frame(path,dataset)
        return thaw(manifest,rows,verify=verify)
    except PyDaraError:raise
    except Exception as exc:
        raise StorageError(f'Import failed ({type(exc).__name__}): {exc}') from exc
